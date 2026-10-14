"""
Advantage Pro exam-centre automation - everything in one program.

  1. Pearson VUE roster + ADV PRO feedback  ->  "Pearson VUE Exam new 2026.xlsx"
       - Chrome (advantagepro94 profile): log in to Pearson VUE Connect,
         Site Manager > Availability > Roster, pick the date, read every
         candidate (name, client, exam code, phone from the "View" screen).
       - Chrome (advpromktg profile): read that day's responses from the
         "ADV PRO FEEDBACK (Responses)" Google Sheet.
       - Back up the Excel file, add/update the rows in the month sheet, save.

  2. LinkedIn from Outlook  ->  Excel "Linkedin Id"
       - Rows with an EmailId but no Linkedin Id.
       - Chrome ("Arish - Chrome" profile): Outlook > New mail, put the email
         in "To", hover -> LinkedIn logo -> "See profile on LinkedIn", take the
         linkedin.com/in/... address. The mail is never sent.
       - Back up the Excel file and write the LinkedIn URLs.

  3. Our Student (Adv Pro)  ->  Excel "Our Student Y/N"
       - Rows with a Phone No but no "Our Student Y/N".
       - Adv Pro.exe (vetri): Master > Registration > Display, Filter By
         PresentPhoneNo, search the phone number: found = "y", not found = "n".

  4. Certifications (LinkedIn)  ->  Excel RHCSA ... Anthropic columns
       - Rows with a Linkedin Id and all certification columns empty.
       - Chrome ("Arish - Chrome" profile): open the profile's
         "Licenses & certifications" list; "y" in every column it shows.

  5. All: 1, 2, 3, 4 for the same date.

  6. Test: offline check of the Excel step with sample data (never saves).

Usage
  python main.py                              # menu
  python main.py pvue    [--date 06-10-2026] [--skip-roster] [--skip-feedback] [--dry-run]
  python main.py linkedin [--date 07-10-2026 | --month 2026-10] [--limit 3] [--ask] [--dry-run]
  python main.py student  [--date 07-10-2026 | --month 2026-10] [--limit 3] [--dry-run]
  python main.py certs    [--date 07-10-2026 | --month 2026-10] [--limit 3] [--dry-run]
  python main.py all     [--date 06-10-2026] [--dry-run]
  python main.py test

  Dates can be typed as DD-MM-YYYY or YYYY-MM-DD.
"""
import argparse
import configparser
import csv
import datetime as dt
import json
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import warnings
from copy import copy
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.ini"
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
UPPER, LOWER = "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)


# ================================================================ helpers

def log(msg):
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


def ask_user(msg):
    input(f"\n>>> {msg}\n    Press Enter when done... ")


def parse_date(text):
    """'08-10-2026', '8/10/2026', '8.10.26', '2026-10-08' or '8-10' (this year). Day comes first."""
    text = text.strip()
    if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", text):
        return dt.date.fromisoformat("-".join(p.zfill(2) for p in text.split("-")))
    parts = re.split(r"[-/.\s]+", text)
    if not 2 <= len(parts) <= 3 or not all(p.isdigit() for p in parts):
        raise ValueError(text)
    day, month = int(parts[0]), int(parts[1])
    year = int(parts[2]) if len(parts) == 3 else dt.date.today().year
    if year < 100:
        year += 2000
    return dt.date(year, month, day)


def ask_date():
    today = dt.date.today()
    while True:
        text = input(f"\n>>> Which date? Type it as DD-MM-YYYY (e.g. {today:%d-%m-%Y}),\n"
                     f"    or just press Enter for today ({today:%d-%b-%Y}): ").strip(" \t﻿")
        if not text:
            return today
        try:
            date = parse_date(text)
        except ValueError:
            print("    Not a valid date - try again, e.g. 07-10-2026")
            continue
        print(f"    Using {date:%A, %d-%b-%Y}")
        return date


def ask_scope():
    """-> (date or None, any day of the month to use)."""
    today = dt.date.today()
    while True:
        text = input(f"\n>>> Which date? Type it as DD-MM-YYYY (e.g. {today:%d-%m-%Y}),\n"
                     f"    or just press Enter for all of {today:%B %Y}: ").strip(" \t﻿")
        if not text:
            print(f"    Using every row of {today:%B %Y}")
            return None, today
        try:
            date = parse_date(text)
        except ValueError:
            print("    Not a valid date - try again, e.g. 07-10-2026")
            continue
        print(f"    Using {date:%A, %d-%b-%Y}")
        return date, date


def load_config():
    cfg = configparser.RawConfigParser(inline_comment_prefixes=(";",))
    cfg.read(CONFIG_PATH, encoding="utf-8")
    return cfg


def save_config_value(section, key, value):
    """Update one key in config.ini without losing the comments."""
    lines = CONFIG_PATH.read_text(encoding="utf-8").splitlines()
    current = None
    for i, line in enumerate(lines):
        s = line.strip()
        if s.startswith("[") and s.endswith("]"):
            current = s[1:-1]
        elif current == section and re.match(rf"{key}\s*=", s):
            lines[i] = f"{key} = {value}"
            break
    CONFIG_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def norm_name(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def clean_cell_text(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def phone10(value):
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else ""


def clean_phone(raw):
    """Indian mobile -> int like the existing sheet; anything else stays as text."""
    raw = clean_cell_text(raw)
    if not raw:
        return None
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 10 and digits[0] in "6789":
        return int(digits)
    return raw


def is_blank(value):
    return value is None or str(value).strip() == ""


def name_tokens(s):
    return re.findall(r"[a-z]+", str(s or "").lower())


def name_score(fb_name, row_name):
    """How well a feedback name ("V.M.Hemanth", "Raghu") fits a roster name
    ("HEMANTH V M", "Raghu Kumar"). 0 = no match, higher = better."""
    f, r = name_tokens(fb_name), name_tokens(row_name)
    if not f or not r or not any(len(t) >= 3 for t in f):
        return 0

    def tok_ok(t):
        if len(t) == 1:  # initial
            return any(x.startswith(t) for x in r)
        return any(t == x or (len(t) >= 3 and len(x) >= 3 and (x.startswith(t) or t.startswith(x)))
                   for x in r)

    if all(tok_ok(t) for t in f):
        return 100 + sum(len(t) for t in f if len(t) >= 3)
    fj, rj = "".join(f), "".join(r)
    if len(fj) >= 5 and (fj in rj or rj in fj):  # "vairamuthu" vs "vaira muthu"
        return 90
    ratio = SequenceMatcher(None, " ".join(sorted(f)), " ".join(sorted(r))).ratio()
    return int(ratio * 80) if ratio >= 0.85 else 0  # small typos: "vimal kimar"


def save_backup(path, backup_dir):
    bdir = HERE / backup_dir
    bdir.mkdir(exist_ok=True)
    backup = bdir / f"{Path(path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    log(f"Backup saved: {backup}")


# ================================================================ chrome

def port_open(port):
    with socket.socket() as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", port)) == 0


def shortcut_profile(name):
    """Read --profile-directory from a Desktop shortcut like 'Advantage pro - Chrome.lnk'."""
    for desktop in (Path.home() / "Desktop", Path(r"C:\Users\Public\Desktop")):
        lnk = desktop / name
        if lnk.exists():
            args = subprocess.run(
                ["powershell", "-NoProfile", "-Command",
                 f"(New-Object -ComObject WScript.Shell).CreateShortcut('{lnk}').Arguments"],
                capture_output=True, text=True).stdout
            m = re.search(r'--profile-directory=(?:"([^"]+)"|(\S+))', args)
            if m:
                return m.group(1) or m.group(2)
            raise RuntimeError(f"No --profile-directory in shortcut {lnk}")
    raise RuntimeError(f"Shortcut not found on the Desktop: {name}")


def chrome_running():
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq chrome.exe"], capture_output=True, text=True).stdout
    return "chrome.exe" in out.lower()


_driver = None

# keep Chrome from freezing/throttling tabs in the background
CHROME_FLAGS = ["--disable-background-timer-throttling", "--disable-renderer-backgrounding",
                "--disable-backgrounding-occluded-windows",
                "--disable-features=IntensiveWakeUpThrottling,CalculateNativeWinOcclusion"]


def chrome_pages(port):
    """Tabs of the debug-port Chrome, straight from Chrome (never hangs on a frozen tab)."""
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/json/list", timeout=10) as r:
            return [t for t in json.load(r) if t.get("type") == "page"]
    except Exception:
        return []


def activate_tab(port, tab_id):
    try:
        urllib.request.urlopen(f"http://127.0.0.1:{port}/json/activate/{tab_id}", timeout=10).close()
    except Exception:
        pass


def close_tab(port, tab_id):
    try:
        urllib.request.urlopen(f"http://127.0.0.1:{port}/json/close/{tab_id}", timeout=10).close()
    except Exception:
        pass


def wake_tabs(port):
    """Chrome freezes tabs left in the background, and ChromeDriver hangs forever
    on a frozen tab (the script then stops on the login page). Bringing each tab
    to the front for a moment wakes it up."""
    for t in chrome_pages(port):
        if t["url"].startswith("http"):
            activate_tab(port, t["id"])
            time.sleep(0.3)


def open_chrome(cfg, section, start_url, match):
    """Open start_url in the Chrome profile of the section's Desktop shortcut and
    return the Selenium driver switched to that window.

    All profiles run in one Chrome, started with a debug port so the script can
    control it. That only works if Chrome was fully closed before the first launch.
    """
    global _driver
    port = cfg.getint("chrome", "debug_port")
    profile = shortcut_profile(cfg.get(section, "shortcut"))
    if not port_open(port) and chrome_running():
        # Chrome often keeps running in the background after its windows are closed
        answer = input("\n>>> Chrome is already running (maybe in the background).\n"
                       "    The script must restart it. Close Chrome now? Open tabs come back\n"
                       "    with Ctrl+Shift+T. [Y/n]: ").strip().lower()
        if answer in ("", "y", "yes"):
            subprocess.run(["taskkill", "/IM", "chrome.exe", "/F"], capture_output=True)
            for _ in range(20):
                if not chrome_running():
                    break
                time.sleep(0.5)
        if chrome_running():
            sys.exit("Chrome is still running - close it (check the system tray) and run again.")
        log("Chrome closed")
    before = {t["id"] for t in chrome_pages(port)}
    log(f"Opening Chrome profile '{profile}' ({cfg.get(section, 'shortcut')})")
    subprocess.Popen([cfg.get("chrome", "path"), f"--profile-directory={profile}",
                      f"--remote-debugging-port={port}", *CHROME_FLAGS, start_url])
    for _ in range(60):
        if port_open(port):
            break
        time.sleep(0.5)
    else:
        raise RuntimeError(f"Chrome did not open the control port {port}")

    # find the tab that just opened start_url (newest first) without touching the others
    target = None
    end = time.time() + 30
    while time.time() < end and not target:
        pages = [t for t in chrome_pages(port) if match in t["url"]]
        new = [t for t in pages if t["id"] not in before]
        target = (new or pages or [None])[0]
        if not target:
            time.sleep(1)
    if not target:
        ask_user(f"Could not find the new Chrome tab ({match}). Click on it in Chrome.")
        target = next((t for t in chrome_pages(port) if match in t["url"]), None)
        if not target:
            raise RuntimeError(f"No Chrome tab with {match}")

    log("Waking up Chrome tabs")
    wake_tabs(port)
    activate_tab(port, target["id"])
    if not _driver:
        opts = Options()
        opts.add_experimental_option("debuggerAddress", f"127.0.0.1:{port}")
        _driver = webdriver.Chrome(options=opts)
        log("Connected to Chrome")
    # ChromeDriver window handle = Chrome tab id. The tab can be replaced while a
    # site redirects (e.g. Outlook sign-in), so look it up again if it's gone.
    for _ in range(3):
        ids = [target["id"]] + [t["id"] for t in chrome_pages(port)
                                if match in t["url"] and t["id"] not in before]
        for tab_id in ids:
            try:
                _driver.switch_to.window(tab_id)
                return _driver
            except Exception:
                pass
        ask_user(f"The {match} tab was closed. Open {start_url} in the right Chrome profile.")
        target = next((t for t in chrome_pages(port) if match in t["url"]), target)
    raise RuntimeError(f"No Chrome tab with {match}")


def visible(el):
    try:
        return el.is_displayed()
    except Exception:
        return False


def click(driver, el):
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def has_text(driver, text):
    xp = f"//*[normalize-space(translate(., '{UPPER}', '{LOWER}'))='{text.lower()}']"
    return any(visible(e) for e in driver.find_elements(By.XPATH, xp))


def click_text(driver, *texts, timeout=10, within=None):
    """Click the visible element whose whole text equals one of `texts` (case-insensitive)."""
    root = within or driver
    end = time.time() + timeout
    while True:
        for text in texts:
            xp = f".//*[normalize-space(translate(., '{UPPER}', '{LOWER}'))='{text.lower()}']"
            els = [e for e in root.find_elements(By.XPATH, xp) if visible(e)]
            if els:
                # prefer real buttons/links over headings with the same text
                clickable = [e for e in els if e.tag_name in ("a", "button")
                             or e.get_attribute("role") in ("button", "link", "menuitem", "tab")
                             or e.find_elements(By.XPATH, "ancestor::a|ancestor::button")]
                click(driver, (clickable or els)[-1])  # last = innermost element
                return True
        if time.time() > end:
            return False
        time.sleep(0.5)


def find_input(driver, keywords):
    for el in driver.find_elements(By.CSS_SELECTOR, "input"):
        if not visible(el):
            continue
        attrs = " ".join((el.get_attribute(a) or "") for a in
                         ("name", "id", "placeholder", "aria-label", "formcontrolname", "type")).lower()
        if any(k in attrs for k in keywords):
            return el
    return None


def type_into(el, value):
    el.click()
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.DELETE)
    el.send_keys(value)


def wait_for(driver, js, *args, timeout=8):
    end = time.time() + timeout
    while True:
        el = driver.execute_script(js, *args)
        if el or time.time() > end:
            return el
        time.sleep(0.5)


# ================================================================ pearson vue

def pearson_login(driver, cfg):
    log("Opening Pearson VUE login page")
    end = time.time() + 30
    pwd = None
    while time.time() < end and not pwd:
        pwd = find_input(driver, ["password"])
        if not pwd:
            if has_text(driver, "Site Manager"):  # already logged in
                log("Already logged in")
                return
            time.sleep(1)
    if not pwd:
        ask_user("Could not find the login form. Please log in manually in the Chrome window.")
        return

    user = find_input(driver, ["user", "login"])
    site = find_input(driver, ["site"])
    if not user or not site:  # fall back to field order: username ... site id
        texts = [e for e in driver.find_elements(By.CSS_SELECTOR, "input")
                 if visible(e) and (e.get_attribute("type") or "text") in ("text", "number", "tel", "email")]
        user = user or (texts[0] if texts else None)
        site = site or (texts[-1] if len(texts) > 1 else None)
    if user:
        type_into(user, cfg.get("pearson", "username"))
    type_into(pwd, cfg.get("pearson", "password"))
    if site:
        type_into(site, cfg.get("pearson", "site_id"))

    if not click_text(driver, "Sign In", "Sign in", "Log In", "Login", "Submit", timeout=3):
        (site or pwd).send_keys(Keys.ENTER)
    log("Login submitted")
    if click_text(driver, "Skip", timeout=20):
        log("Clicked Skip")


def open_roster(driver, date, cfg):
    for step in ("Site Manager", "Availability", "Roster"):
        if click_text(driver, step, timeout=25):
            log(f"Clicked {step}")
            time.sleep(2)
        else:
            ask_user(f"Could not find '{step}'. Please click it in the Pearson VUE window.")
    set_roster_date(driver, date)
    time.sleep(3)
    if cfg.get("pearson", "confirm_roster", fallback="yes").lower().startswith("y"):
        ask_user(f"Check the roster is showing {date:%d-%b-%Y} (fix it in Chrome if not).")


def set_roster_date(driver, date):
    """Best effort: fill the date box in the format it already uses, then Search."""
    today = dt.date.today()
    for el in driver.find_elements(By.CSS_SELECTOR, "input"):
        if not visible(el):
            continue
        attrs = " ".join((el.get_attribute(a) or "") for a in
                         ("type", "name", "id", "placeholder", "aria-label", "formcontrolname")).lower()
        if "date" not in attrs:
            continue
        if (el.get_attribute("type") or "") == "date":
            driver.execute_script(
                "arguments[0].value=arguments[1];"
                "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                el, date.isoformat())
        else:
            current = el.get_attribute("value") or ""
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", current):
                text = date.strftime("%Y-%m-%d")
            elif current == today.strftime("%d/%m/%Y") and today.day != today.month:
                text = date.strftime("%d/%m/%Y")
            else:
                text = date.strftime("%m/%d/%Y")
            type_into(el, text)
            el.send_keys(Keys.TAB)
        log(f"Set roster date {date}")
        click_text(driver, "Search", "Go", "Apply", "View", timeout=3)
        return True
    if date != today:
        ask_user(f"Could not find the date box. Please pick {date:%d-%b-%Y} on the roster.")
    return False


GRAB_TABLE_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const norm = s => (s || '').replace(/\s+/g, ' ').trim();
const groups = [];
for (const t of document.querySelectorAll('table')) if (vis(t)) groups.push([...t.querySelectorAll('tr')]);
const rr = [...document.querySelectorAll('[role=row]')];
if (rr.length) groups.push(rr);
for (const rows of groups) {
  const data = rows.filter(vis).map(r => ({el: r, cells: [...r.querySelectorAll(
      'th,td,[role=columnheader],[role=gridcell],[role=cell]')].map(c => norm(c.innerText))}));
  const hi = data.findIndex(d => d.cells.some(c => /last\s*name/i.test(c)));
  if (hi >= 0) return {header: data[hi].cells,
                       rows: data.slice(hi + 1).filter(d => d.cells.length >= 3).map(d => [d.el, d.cells])};
}
return null;
"""


def grab_roster(driver, timeout=20, min_rows=0):
    end = time.time() + timeout
    while True:
        table = driver.execute_script(GRAB_TABLE_JS)
        if (table and len(table["rows"]) >= min_rows) or time.time() > end:
            return table
        time.sleep(1)


def roster_columns(header):
    h = [c.lower() for c in header]

    def idx(*keys):
        for k in keys:
            for i, c in enumerate(h):
                if k in c:
                    return i
        return None

    return {"first": idx("first"), "last": idx("last name", "last"),
            "client": idx("client"), "code": idx("exam code", "exam")}


def view_text(driver):
    return driver.execute_script(r"""
        const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
        const d = [...document.querySelectorAll('[role=dialog],.modal,.modal-dialog,mat-dialog-container')].filter(vis);
        return (d.length ? d[d.length - 1] : document.body).innerText;""") or ""


def phone_in_text(text):
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for i, line in enumerate(lines):
        if re.search(r"phone|mobile|contact number|telephone", line, re.I):
            for cand in [line] + lines[i + 1:i + 3]:
                m = re.search(r"\+?\d[\d\s\-().]{6,}\d", cand)
                if m:
                    return m.group(0).strip()
    return ""


def read_phone_from_view(driver, name, timeout=12):
    """Wait for the View screen to show a phone number; if it never does, save
    the screen text in logs/ so we can see what the page looked like."""
    end = time.time() + timeout
    text = ""
    while time.time() < end:
        text = view_text(driver)
        phone = phone_in_text(text)
        if phone:
            return phone
        time.sleep(1)
    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    dump = logs / f"view_{re.sub(r'[^A-Za-z0-9]+', '_', name)}.txt"
    dump.write_text(f"URL: {driver.current_url}\n\n{text}", encoding="utf-8")
    log(f"  (View screen text saved to {dump.name})")
    return ""


def close_view(driver):
    has_dialog = driver.execute_script(
        "return [...document.querySelectorAll('[role=dialog],.modal,mat-dialog-container')]"
        ".some(e => e.offsetWidth || e.offsetHeight);")
    if has_dialog and click_text(driver, "Close", "Cancel", "OK", "Back", "×", timeout=3):
        return
    if not has_dialog and click_text(driver, "Back", "Back to roster", "Return to roster", timeout=2):
        return
    if has_dialog:
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
    else:
        driver.back()


def back_to_roster(driver, date):
    """After reading a phone number: click the 'Roster' button to return to the list."""
    if click_text(driver, "Roster", timeout=5):
        time.sleep(2)
        if date != dt.date.today():  # roster may reset to today
            set_roster_date(driver, date)
            time.sleep(2)
    else:
        log("  'Roster' button not found - closing the view instead")
        close_view(driver)
        time.sleep(1.5)


def scrape_roster(driver, cfg, date):
    pearson_login(driver, cfg)
    open_roster(driver, date, cfg)

    table = grab_roster(driver)
    if not table:
        ask_user("Could not read the roster table. Make sure the roster list is visible.")
        table = grab_roster(driver)
        if not table:
            raise RuntimeError("Roster table not found")
    cols = roster_columns(table["header"])
    log(f"Roster columns: {table['header']}")

    def cell(cells, key):
        i = cols[key]
        return cells[i] if i is not None and i < len(cells) else ""

    people = []
    for _, cells in table["rows"]:
        name = clean_cell_text(f"{cell(cells, 'first')} {cell(cells, 'last')}")
        if name:
            people.append({"name": name, "client": cell(cells, "client"),
                           "code": cell(cells, "code"), "phone": ""})
    log(f"{len(people)} exam(s) on the roster")

    for i, person in enumerate(people):
        table = grab_roster(driver, min_rows=i + 1)
        if not table or len(table["rows"]) <= i:  # page reset after Back -> open roster again
            log("  roster not showing - opening it again")
            open_roster(driver, date, cfg)
            table = grab_roster(driver, min_rows=i + 1)
        if not table or i >= len(table["rows"]):
            log(f"  ! could not find row for {person['name']}, phone left blank")
            continue
        row_el, cells = table["rows"][i]
        if norm_name(cell(cells, "first")) not in norm_name(person["name"]):
            log(f"  ! row order changed at {person['name']}, phone left blank")
            continue
        if not click_text(driver, "View", timeout=5, within=row_el):
            log(f"  - {person['name']}: no View button")
            continue
        time.sleep(1)
        person["phone"] = read_phone_from_view(driver, person["name"])
        log(f"  {person['name']}: {person['phone'] or '(no phone found)'}")
        back_to_roster(driver, date)
    return people


# ================================================================ google sheet feedback

def find_feedback_sheet(driver, cfg):
    url = cfg.get("feedback", "sheet_url", fallback="").strip()
    if url:
        return url
    driver.get("https://drive.google.com/drive/search?q=ADV%20PRO%20FEEDBACK%20(Responses)")
    time.sleep(3)
    if "accounts.google.com" in driver.current_url:
        ask_user("Sign in to advpromktg@gmail.com in the Chrome window (only needed once).")
        driver.get("https://drive.google.com/drive/search?q=ADV%20PRO%20FEEDBACK%20(Responses)")
    ask_user("Double-click 'ADV PRO FEEDBACK (Responses)' in Google Drive so the sheet opens.")
    port = cfg.getint("chrome", "debug_port")
    for tab in chrome_pages(port):
        if "/spreadsheets/d/" in tab["url"]:
            activate_tab(port, tab["id"])
            driver.switch_to.window(tab["id"])
            url = tab["url"]
            save_config_value("feedback", "sheet_url", url)
            log("Saved the feedback sheet link in config.ini")
            return url
    raise RuntimeError("The feedback sheet was not opened")


def download_feedback_csv(driver, sheet_url):
    m = re.search(r"/spreadsheets/d/([\w-]+)", sheet_url)
    gid = re.search(r"[#&?]gid=(\d+)", sheet_url)
    export = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    if gid:
        export += f"&gid={gid.group(1)}"
    dl = HERE / "downloads"
    dl.mkdir(exist_ok=True)
    for f in dl.glob("*.csv"):
        f.unlink()
    try:
        driver.execute_cdp_cmd("Browser.setDownloadBehavior", {"behavior": "allow", "downloadPath": str(dl)})
    except Exception:
        driver.execute_cdp_cmd("Page.setDownloadBehavior", {"behavior": "allow", "downloadPath": str(dl)})
    driver.get(export)
    for _ in range(60):
        files = [f for f in dl.glob("*.csv")]
        if files and not list(dl.glob("*.crdownload")):
            return files[0]
        time.sleep(1)
    raise RuntimeError("Feedback CSV did not download (are you signed in to advpromktg@gmail.com?)")


def timestamp_dates(text, formats=None):
    text = text.strip()
    formats = formats or ["%m/%d/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                          "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"]
    out = set()
    for f in formats:
        try:
            out.add(dt.datetime.strptime(text, f).date())
        except ValueError:
            pass
    return out


def detect_ts_format(values):
    """Google exports dates in the sheet's locale: decide month/day vs day/month
    from values like 10/27/2026 or 27/10/2026, so 7/10/2026 isn't read both ways."""
    mdy = dmy = 0
    for v in values:
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})/\d{4}", v)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            mdy += b > 12
            dmy += a > 12
    if mdy == dmy:
        return ""
    return "%m/%d/%Y" if mdy > dmy else "%d/%m/%Y"


def read_feedback(csv_path, date, ts_format):
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return []
    header = [h.lower() for h in rows[0]]

    def col(test):
        return next((i for i, h in enumerate(header) if test(h)), None)

    c = {"ts": col(lambda h: "timestamp" in h),
         "phone": col(lambda h: "phone" in h or "mobile" in h or "contact" in h),
         "email": col(lambda h: "email" in h),
         "company": col(lambda h: "company" in h),
         "designation": col(lambda h: "designation" in h),
         "experience": col(lambda h: "experience" in h),
         "name": col(lambda h: "name" in h and "company" not in h)}
    log(f"Feedback sheet columns found: { {k: rows[0][v] for k, v in c.items() if v is not None} }")

    def get(r, key):
        i = c[key]
        return r[i].strip() if i is not None and i < len(r) else ""

    formats = [ts_format] if ts_format else None
    if not formats:
        order = detect_ts_format(get(r, "ts") for r in rows[1:])
        if order:  # with or without the time
            formats = [order + " %H:%M:%S", order]
            log(f"Feedback timestamps are {'month/day' if order.startswith('%m') else 'day/month'}")
    out = []
    for r in rows[1:]:
        if date in timestamp_dates(get(r, "ts"), formats):
            out.append({k: get(r, k) for k in c})
    log(f"{len(out)} feedback response(s) on {date}")
    return out


def fetch_feedback(cfg, date):
    url = cfg.get("feedback", "sheet_url", fallback="").strip() or "https://drive.google.com/"
    driver = open_chrome(cfg, "feedback", url, "google.com")
    if "accounts.google.com" in driver.current_url:
        ask_user("Sign in to Google in the Ranjit Chrome window.")
    url = find_feedback_sheet(driver, cfg)
    csv_path = download_feedback_csv(driver, url)
    return read_feedback(csv_path, date, cfg.get("feedback", "timestamp_format", fallback="").strip())


# ================================================================ excel

def check_excel_closed(path):
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        sys.exit(f"The Excel file is open on some computer. Close it and run again:\n  {path}")


def month_sheet(wb, date):
    abbr = MONTHS[date.month - 1]
    for ws in wb.worksheets:
        if ws.title.strip().upper().startswith(abbr):
            return ws
    # New month: copy the previous month's sheet and keep only the header
    prev = [ws for ws in wb.worksheets if ws.title.strip().upper()[:3] in MONTHS][-1]
    ws = wb.copy_worksheet(prev)
    ws.title = abbr
    hr = header_row(ws)
    if ws.max_row > hr:
        ws.delete_rows(hr + 1, ws.max_row - hr)
    if isinstance(ws["A1"].value, str):
        ws["A1"].value = re.sub(r"for\s+.*$", f"for  {date:%B %Y}", ws["A1"].value)
    log(f"Created new sheet {abbr}")
    return ws


def header_row(ws):
    for r in range(1, 20):
        if str(ws.cell(r, 1).value or "").strip().lower() == "date":
            return r
    raise RuntimeError(f"Header row ('Date' in column A) not found in sheet {ws.title}")


def sheet_columns(ws, hr):
    cols, after_feedback = {}, False
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(hr, c).value or "").strip().lower()
        if not h:
            continue
        cols.setdefault(h, c)
        if h == "adv pro feedback":
            after_feedback = True
        elif h == "remark" and after_feedback and "feedback remark" not in cols:
            cols["feedback remark"] = c
    for need in ("date", "name", "phone no", "emailid", "client", "exam code",
                 "adv pro feedback", "feedback remark", "company name", "designation"):
        if need not in cols:
            raise RuntimeError(f"Column '{need}' not found in sheet {ws.title}")
    return cols


def linkedin_col(ws, hr):
    for c in range(1, ws.max_column + 1):
        if re.sub(r"\s+", "", str(ws.cell(hr, c).value or "")).lower() == "linkedinid":
            return c
    raise RuntimeError(f"Column 'Linkedin Id' not found in sheet {ws.title}")


def last_data_row(ws, hr):
    for r in range(ws.max_row, hr, -1):
        if not is_blank(ws.cell(r, 1).value) or not is_blank(ws.cell(r, 2).value):
            return r
    return hr


def cell_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def update_excel(path, date, people, feedback, backup_dir, dry_run=False):
    wb = openpyxl.load_workbook(path)
    ws = month_sheet(wb, date)
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)
    last = last_data_row(ws, hr)

    # --- roster: add people not already on the sheet for this date
    existing = {norm_name(ws.cell(r, cols["name"]).value)
                for r in range(hr + 1, last + 1) if cell_date(ws.cell(r, cols["date"]).value) == date}
    added = 0
    for p in people:
        if norm_name(p["name"]) in existing:
            log(f"  already in sheet: {p['name']}")
            continue
        last += 1
        style_row = last - 1 if last - 1 > hr else None
        for c in range(1, ws.max_column + 1):
            if style_row:
                ws.cell(last, c)._style = copy(ws.cell(style_row, c)._style)
        ws.cell(last, cols["date"]).value = dt.datetime(date.year, date.month, date.day)
        ws.cell(last, cols["name"]).value = p["name"]
        ws.cell(last, cols["phone no"]).value = clean_phone(p["phone"])
        ws.cell(last, cols["client"]).value = p["client"]
        ws.cell(last, cols["exam code"]).value = p["code"]
        existing.add(norm_name(p["name"]))
        added += 1
        log(f"  added: {p['name']} | {p['client']} | {p['code']} | {p['phone']}")

    # --- feedback: match by phone or email (that day first, then the month), else by name (that day)
    updated, unmatched = 0, []
    day_rows = [r for r in range(hr + 1, last + 1) if cell_date(ws.cell(r, cols["date"]).value) == date]
    month_rows = list(range(hr + 1, last + 1))

    def contact_match(fb, rows):
        p = phone10(fb["phone"])
        email = fb["email"].strip().lower()
        for r in rows:
            if p and p == phone10(ws.cell(r, cols["phone no"]).value):
                return r, "phone"
            if email and email == str(ws.cell(r, cols["emailid"]).value or "").strip().lower():
                return r, "email"
        return None, ""

    def name_match(fb):
        scored = sorted(((name_score(fb["name"], ws.cell(r, cols["name"]).value), r) for r in day_rows),
                        reverse=True)
        if not scored or scored[0][0] == 0:
            return None, ""
        if len(scored) > 1 and scored[1][0] == scored[0][0]:
            names = [clean_cell_text(ws.cell(r, cols["name"]).value) for s, r in scored if s == scored[0][0]]
            log(f"  ? '{fb['name']}' fits more than one row ({', '.join(names)}) - fill it by hand")
            return None, ""
        return scored[0][1], "name"

    for fb in feedback:
        row, how = contact_match(fb, day_rows)
        if not row:
            row, how = name_match(fb)
        if not row:
            row, how = contact_match(fb, month_rows)
        if not row:
            unmatched.append(fb)
            continue

        def fill(col, value, overwrite=False):
            if value and (overwrite or is_blank(ws.cell(row, cols[col]).value)):
                ws.cell(row, cols[col]).value = value

        # the phone the candidate typed in the Google form wins over the Pearson one
        fill("phone no", clean_phone(fb["phone"]), overwrite=True)
        fill("emailid", fb["email"])
        fill("company name", clean_cell_text(fb["company"]))
        fill("designation", clean_cell_text(fb["designation"]))
        fill("feedback remark", clean_cell_text(fb["experience"]))
        fill("adv pro feedback", "y")
        updated += 1
        log(f"  feedback '{fb['name']}' -> row {row}: {clean_cell_text(ws.cell(row, cols['name']).value)}"
            f" (matched by {how}) | phone {fb['phone']} | {fb['email']}")

    for fb in unmatched:
        log(f"  ! feedback not matched to any row (not on the roster?): "
            f"{fb['name'] or ''} | {fb['phone']} | {fb['email']}")
    log(f"Sheet {ws.title}: {added} row(s) added, {updated} row(s) updated with feedback")

    if dry_run:
        log("Dry run - Excel file NOT saved")
        return
    if not added and not updated:
        log("Nothing changed - Excel file not saved")
        return
    save_backup(path, backup_dir)
    wb.save(path)
    log("Excel file saved")


def linkedin_rows_to_do(path, date, month_day):
    wb = openpyxl.load_workbook(path)
    ws = month_sheet(wb, month_day)
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)
    li = linkedin_col(ws, hr)
    todo = []
    for r in range(hr + 1, last_data_row(ws, hr) + 1):
        if date and cell_date(ws.cell(r, cols["date"]).value) != date:
            continue
        email = str(ws.cell(r, cols["emailid"]).value or "").strip()
        if not EMAIL_RE.match(email) or not is_blank(ws.cell(r, li).value):
            continue
        todo.append({"row": r, "name": clean_cell_text(ws.cell(r, cols["name"]).value),
                     "email": email, "linkedin": ""})
    return ws.title, todo


def save_linkedin_results(path, sheet, results, backup_dir):
    found = [x for x in results if x["linkedin"]]
    if not found:
        log("No LinkedIn URLs found - Excel file not changed")
        return
    while True:
        try:
            with open(path, "r+b"):
                break
        except PermissionError:
            ask_user(f"The Excel file is open somewhere. Close it so the {len(found)} LinkedIn URL(s) can be saved.")
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)
    li = linkedin_col(ws, hr)
    written = 0
    for x in found:
        row = x["row"]
        # the row may have moved if someone edited the file meanwhile -> find it by email
        if str(ws.cell(row, cols["emailid"]).value or "").strip().lower() != x["email"].lower():
            row = next((r for r in range(hr + 1, ws.max_row + 1)
                        if str(ws.cell(r, cols["emailid"]).value or "").strip().lower() == x["email"].lower()), None)
        if not row or not is_blank(ws.cell(row, li).value):
            continue
        cell = ws.cell(row, li)
        above = ws.cell(row - 1, li)
        if row - 1 > hr and above.has_style:
            cell._style = copy(above._style)
        cell.value = x["linkedin"]
        written += 1
        log(f"  row {row}: {x['name']} -> {x['linkedin']}")
    if not written:
        log("Nothing to write - Excel file not changed")
        return
    save_backup(path, backup_dir)
    wb.save(path)
    log(f"Excel file saved - {written} LinkedIn URL(s) written to sheet {sheet}")


# ================================================================ outlook -> linkedin

TO_BOX_CSS = ['div[contenteditable="true"][aria-label="To"]', 'input[aria-label="To"]',
              '[role="textbox"][aria-label^="To"]', 'div[contenteditable="true"][aria-label*="To"]']


def find_to_box(driver):
    for css in TO_BOX_CSS:
        for el in driver.find_elements(By.CSS_SELECTOR, css):
            if visible(el):
                return el
    return None


def open_compose(driver):
    end = time.time() + 40
    while time.time() < end:
        if find_to_box(driver):
            return find_to_box(driver)
        if "login" in driver.current_url or "signin" in driver.current_url:
            ask_user("Sign in to Outlook in the Arish Chrome window.")
            end = time.time() + 40
        elif click_text(driver, "New mail", "New message", "New email", timeout=3):
            log("Clicked New mail")
            time.sleep(3)
        else:
            time.sleep(1)
    ask_user("Could not open a new mail. Click 'New mail' in Outlook.")
    box = find_to_box(driver)
    if not box:
        raise RuntimeError("The 'To' box of the new mail was not found")
    return box


FIND_PILL_JS = r"""
const email = arguments[0].toLowerCase(), box = arguments[1];
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const has = e => [e.getAttribute('aria-label'), e.getAttribute('title'), e.innerText]
                 .some(v => (v || '').toLowerCase().includes(email));
// the recipient "pill" sits next to the To box - search the compose header around it
let area = box;
for (let i = 0; i < 6 && area.parentElement; i++) area = area.parentElement;
const hits = [...area.querySelectorAll('*')].filter(e => vis(e) && e !== box && has(e) && !e.isContentEditable);
// innermost element that matches
return hits.filter(e => !hits.some(o => o !== e && e.contains(o)))[0] || null;
"""

FIND_LINKEDIN_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const label = e => [e.getAttribute('aria-label'), e.getAttribute('title'), e.getAttribute('alt'),
                    e.getAttribute('data-icon-name'), e.innerText].join(' ');
const els = [...document.querySelectorAll('button,a,[role=button],[role=tab],[role=link],img,svg,i')]
            .filter(e => vis(e) && /linked\s*in/i.test(label(e)) && !/see\s+profile|view\s+(full\s+)?profile/i.test(label(e)));
if (!els.length) return null;
const e = els[0];
return e.closest('button,a,[role=button],[role=tab],[role=link]') || e;
"""

FIND_SEE_PROFILE_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const els = [...document.querySelectorAll('a,button,[role=button],[role=link]')].filter(e => vis(e) &&
    /(see|view)\s+(full\s+)?profile\s+on\s+linked\s*in|profile\s+on\s+linked\s*in/i.test(
        (e.innerText || '') + ' ' + (e.getAttribute('aria-label') || '') + ' ' + (e.getAttribute('title') || '')));
return els[0] || null;
"""


def linkedin_profile(url):
    """Any URL (also LinkedIn redirect / sign-in links) -> https://www.linkedin.com/in/<id>/"""
    for _ in range(3):
        m = re.search(r"linkedin\.com/in/([^/?#&\s\"']+)", url or "", re.I)
        if m:
            return f"https://www.linkedin.com/in/{m.group(1)}/"
        url = urllib.parse.unquote(url or "")
    return ""


def hover(driver, el):
    """Put the mouse on the element, then move it slightly (the card shows on a light move)."""
    ac = ActionChains(driver)
    ac.move_to_element(el).pause(0.4).move_by_offset(3, 0).pause(0.4).move_by_offset(-3, 1).perform()


def new_linkedin_tab(port, before, timeout=20):
    """Wait for a tab that wasn't open before to show a LinkedIn address; return its URL and close it."""
    end = time.time() + timeout
    seen = None
    while time.time() < end:
        for t in chrome_pages(port):
            if t["id"] in before:
                continue
            seen = t
            if linkedin_profile(t["url"]):
                close_tab(port, t["id"])
                return t["url"]
        time.sleep(0.5)
    if seen:
        log(f"    new tab did not reach a profile: {seen['url'][:100]}")
        close_tab(port, seen["id"])
    return ""


def clear_to(driver):
    box = find_to_box(driver)
    if not box:
        return
    click(driver, box)
    for _ in range(4):  # Backspace in an empty To box removes the recipient before it
        box.send_keys(Keys.BACKSPACE)
        time.sleep(0.2)


def close_card(driver):
    box = find_to_box(driver)
    if box:
        ActionChains(driver).move_to_element(box).perform()  # mouse away from the pill
    if wait_for(driver, FIND_LINKEDIN_JS, timeout=0) or wait_for(driver, FIND_SEE_PROFILE_JS, timeout=0):
        ActionChains(driver).send_keys(Keys.ESCAPE).perform()
        time.sleep(0.5)


def lookup(driver, port, outlook_tab, email):
    """Email -> LinkedIn profile URL ('' if Outlook shows none)."""
    box = find_to_box(driver) or open_compose(driver)
    clear_to(driver)
    box = find_to_box(driver)
    click(driver, box)
    box.send_keys(email)
    time.sleep(1.5)
    box.send_keys(";")  # turns the typed address into a recipient
    pill = wait_for(driver, FIND_PILL_JS, email, box, timeout=8)
    if not pill:
        log("    recipient not found in the To box")
        return ""

    hover(driver, pill)
    li = wait_for(driver, FIND_LINKEDIN_JS, timeout=6)
    if not li:  # hover again, then try a click on the recipient
        hover(driver, pill)
        li = wait_for(driver, FIND_LINKEDIN_JS, timeout=4)
    if not li:
        click(driver, pill)
        li = wait_for(driver, FIND_LINKEDIN_JS, timeout=6)
    if not li:
        log("    contact card has no LinkedIn logo")
        return ""
    click(driver, li)

    see = wait_for(driver, FIND_SEE_PROFILE_JS, timeout=10)
    if not see:
        log("    no 'See profile on LinkedIn' (no LinkedIn match for this email?)")
        return ""
    direct = linkedin_profile(see.get_attribute("href") or "")
    before = {t["id"] for t in chrome_pages(port)}
    click(driver, see)
    url = new_linkedin_tab(port, before) or direct
    driver.switch_to.window(outlook_tab)
    return linkedin_profile(url)


def manual_lookup(driver, port, outlook_tab, person):
    before = {t["id"] for t in chrome_pages(port)}
    ask_user(f"Open the LinkedIn profile of {person['name']} ({person['email']}) by hand "
             "(hover -> LinkedIn logo -> See profile on LinkedIn).\n"
             "    Leave the LinkedIn tab open. Just press Enter to skip.")
    url = ""
    for t in chrome_pages(port):
        if t["id"] not in before and linkedin_profile(t["url"]):
            url = t["url"]
            close_tab(port, t["id"])
    driver.switch_to.window(outlook_tab)
    return linkedin_profile(url)


def discard_draft(driver):
    clear_to(driver)
    # the Discard button is an icon - find it by its label
    btn = next((e for e in driver.find_elements(By.CSS_SELECTOR, '[aria-label="Discard"]') if visible(e)), None)
    if btn:
        click(driver, btn)
    elif not click_text(driver, "Discard", timeout=3):
        log("Discard button not found - delete the empty draft from Outlook Drafts")
        return
    time.sleep(1)
    click_text(driver, "OK", "Discard", "Yes", timeout=3)  # confirmation, if Outlook asks
    log("Draft discarded")


# ================================================================ excel: rows to check / write back

def month_rows(path, date, month_day, need):
    """Rows of the day (or the whole month when date is None).
    -> (sheet title, {column name: col}, [row numbers]). `need` = columns that must exist."""
    wb = openpyxl.load_workbook(path)
    ws = month_sheet(wb, month_day)
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)
    missing = [n for n in need if n not in cols]
    if missing:
        raise RuntimeError(f"Column(s) {', '.join(missing)} not found in sheet {ws.title}")
    rows = [r for r in range(hr + 1, last_data_row(ws, hr) + 1)
            if not date or cell_date(ws.cell(r, cols["date"]).value) == date]
    return ws, cols, rows


def key_value(col, value):
    return phone10(value) if col == "phone no" else str(value or "").strip().lower().rstrip("/")


def save_cells(path, sheet, updates, backup_dir, what):
    """updates: [{"row", "key": (column name, value), "values": {column name: value}, "label"}].
    The key finds the row again if someone moved it meanwhile. Only empty cells are filled."""
    if not updates:
        log(f"No {what} to write - Excel file not changed")
        return
    while True:
        try:
            with open(path, "r+b"):
                break
        except PermissionError:
            ask_user(f"The Excel file is open somewhere. Close it so the {what} can be saved.")
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)
    written = 0
    for u in updates:
        key_col, key = u["key"]

        def same(r):
            return key_value(key_col, ws.cell(r, cols[key_col]).value) == key_value(key_col, key)

        row = u["row"] if same(u["row"]) else next((r for r in range(hr + 1, ws.max_row + 1) if same(r)), None)
        if not row:
            log(f"  ! row of {u['label']} not found any more - skipped")
            continue
        changed = False
        for name, value in u["values"].items():
            cell = ws.cell(row, cols[name])
            if not is_blank(cell.value):
                continue
            above = ws.cell(row - 1, cols[name])
            if row - 1 > hr and above.has_style:
                cell._style = copy(above._style)
            cell.value = value
            changed = True
        if changed:
            written += 1
            log(f"  row {row}: {u['label']}")
    if not written:
        log(f"No {what} to write - Excel file not changed")
        return
    save_backup(path, backup_dir)
    wb.save(path)
    log(f"Excel file saved - {what} written for {written} row(s) in sheet {sheet}")


# ================================================================ adv pro -> our student

class AdvPro:
    """Drives the Adv Pro desktop program: login, Master > Registration > Display,
    Filter By PresentPhoneNo, and tells if a phone number is registered."""

    GRID = "DataGridWndClass"

    def __init__(self, cfg):
        from pywinauto import Application
        # pywinauto turns its own warnings back on when imported
        warnings.filterwarnings("ignore", message=".*32-bit application.*")
        self.cfg = cfg
        exe = cfg.get("advpro", "exe")
        try:
            self.app = Application(backend="win32").connect(path=Path(exe).name, timeout=2)
            log("Adv Pro is already open")
        except Exception:
            log("Opening Adv Pro")
            self.app = Application(backend="win32").start(f'"{exe}"', work_dir=str(Path(exe).parent))
        self.main = self.app.window(class_name="ThunderRT6MDIForm")
        self.login()
        self.form = self.open_display()
        self.grid = self.form.child_window(class_name=self.GRID).wrapper_object()
        self.search = self.form.child_window(class_name="ThunderRT6TextBox", control_id=51).wrapper_object()
        combo = self.form.child_window(class_name="ThunderRT6ComboBox", control_id=52).wrapper_object()
        combo.select("PresentPhoneNo")
        time.sleep(0.5)
        if combo.selected_text() != "PresentPhoneNo":
            ask_user("Please choose 'PresentPhoneNo' in 'Filter By' in Adv Pro.")
        log("Adv Pro: Filter By = PresentPhoneNo")

    def login(self):
        login = self.app.window(title="Form1", class_name="ThunderRT6FormDC")
        end = time.time() + 30
        while time.time() < end:
            if self.main.exists() and self.main.is_visible():
                return
            if login.exists() and login.is_visible():
                break
            time.sleep(1)
        else:
            ask_user("Adv Pro did not show the login screen. Log in by hand.")
            return
        user = login.child_window(class_name="ThunderRT6TextBox", control_id=1).wrapper_object()
        pwd = login.child_window(class_name="ThunderRT6TextBox", control_id=2).wrapper_object()
        self.set_text(user, self.cfg.get("advpro", "username"))
        self.set_text(pwd, self.cfg.get("advpro", "password"))
        pwd.set_focus()
        pwd.type_keys("{ENTER}")
        if not self.main.wait("visible", timeout=30):
            ask_user("Adv Pro login did not finish. Log in by hand.")
        log("Adv Pro: logged in")

    def display_form(self):
        for child in self.main.children(class_name="ThunderRT6FormDC"):
            if child.is_visible() and child.children(class_name=self.GRID):
                return self.main.child_window(handle=child.handle)
        return None

    def open_display(self):
        form = self.display_form()
        if not form:
            self.main.set_focus()
            self.main.menu_select("Master->Registration->Display")
            log("Adv Pro: Master > Registration > Display")
            end = time.time() + 20
            while not form and time.time() < end:
                time.sleep(1)
                form = self.display_form()
        if not form:
            ask_user("Open Master > Registration > Display in Adv Pro.")
            form = self.display_form()
        if not form:
            raise RuntimeError("Adv Pro Registration Display screen not found")
        return form

    @staticmethod
    def set_text(ctrl, text):
        # VB6 text boxes ignore Ctrl+A, so set the text directly
        import win32con
        import win32gui
        win32gui.SendMessage(ctrl.handle, win32con.WM_SETTEXT, 0, text)

    def grid_has_rows(self):
        """The list shows columns (and a scroll bar) only when something was found.
        Read it two ways: UI Automation, and the screen pixels of the list."""
        from pywinauto.controls.uiawrapper import UIAWrapper
        from pywinauto.uia_element_info import UIAElementInfo
        try:
            uia = len(UIAWrapper(UIAElementInfo(self.grid.handle)).children()) > 0
        except Exception:
            uia = None
        try:
            from PIL import ImageGrab
            r = self.grid.rectangle()
            img = ImageGrab.grab(bbox=(r.left, r.top, r.right, r.bottom), all_screens=True).convert("RGB")
            w, h = img.size
            px = img.load()
            # an empty list is grey right of the first two blank columns
            white = sum(1 for x in range(260, w - 4, 2) for y in range(2, h - 18, 2) if min(px[x, y]) > 245)
            pixels = white > 150
        except Exception:
            pixels = None
        return uia, pixels

    def is_student(self, phone):
        p = phone10(phone)
        if not p:
            return ""
        self.set_text(self.search, p)
        time.sleep(self.cfg.getfloat("advpro", "search_wait", fallback=2))
        last = self.grid_has_rows()
        for _ in range(6):  # wait until the list stops changing
            time.sleep(0.7)
            now = self.grid_has_rows()
            if now == last:
                break
            last = now
        uia, pixels = last
        if uia is not None and pixels is not None and uia != pixels:
            log(f"  ? Adv Pro {p}: unsure (list check {uia}, screen check {pixels}) - check by hand")
        found = uia if uia is not None else pixels
        return "y" if found else "n"


# ================================================================ linkedin -> certifications

# Excel column (lower case) -> what to look for in the "Licenses & certifications" list
CERTS = {
    "rhcsa": r"\bRHCSA\b|Red Hat Certified System Administrator|\bEX200\b",
    "rhce": r"\bRHCE\b|Red Hat Certified Engineer|\bEX294\b",
    "ex280": r"\bEX280\b|OpenShift Administration",
    "ex374": r"\bEX374\b|Developing Automation with Ansible",
    "aws": r"\bAWS\b|Amazon Web Services",
    "azure": r"\bAzure\b|\bAZ-\d{3}\b",
    "cka": r"\bCKA\b|Certified Kubernetes Administrator",
    "ccna": r"\bCCNA\b|Cisco Certified Network Associate",
    "vmware": r"VMware|\bVCP\b",
    "python": r"\bPython\b|\bPCEP\b|\bPCAP\b|\bPCPP\b",
    "anthropic": r"Anthropic|Claude Certified",
}
LINKEDIN_LOGIN = ("/login", "/authwall", "/checkpoint", "/uas/", "/signup")


def main_text(driver):
    return driver.execute_script(
        "const m = document.querySelector('main'); return (m || document.body).innerText;") or ""


def read_certifications(driver, profile_url):
    """Text of the profile's 'Licenses & certifications' list.
    '' = the profile has none, None = the page could not be opened."""
    url = profile_url.rstrip("/") + "/details/certifications/"
    for _ in range(3):
        driver.get(url)
        time.sleep(3)
        if any(k in driver.current_url for k in LINKEDIN_LOGIN):
            ask_user("Sign in to LinkedIn in the Arish Chrome window.")
            continue
        end = time.time() + 20
        while time.time() < end:
            if "/details/certifications" not in driver.current_url:
                return ""  # LinkedIn sends you back to the profile when there are none
            if re.search(r"certifications|nothing to see", main_text(driver), re.I):
                break
            time.sleep(1)
        for _ in range(4):  # the list loads more while scrolling
            driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(1.2)
            click_text(driver, "Show more results", timeout=0)
        text = main_text(driver)
        if re.search(r"nothing to see", text, re.I):
            return ""
        return text
    return None


def find_certs(text):
    return [col for col, pattern in CERTS.items() if re.search(pattern, text or "", re.I)]


# ================================================================ tasks

def run_pvue(cfg, date, args):
    """Pearson VUE roster + Google Sheet feedback -> Excel."""
    excel = cfg.get("excel", "path")
    log(f"=== Pearson VUE roster + feedback | Date: {date:%d-%b-%Y}   Excel: {excel}")
    check_excel_closed(excel)

    people, feedback = [], []
    if not args.skip_roster:
        driver = open_chrome(cfg, "pearson", cfg.get("pearson", "url"), "pearsonvue.com")
        people = scrape_roster(driver, cfg, date)
    if not args.skip_feedback:
        feedback = fetch_feedback(cfg, date)

    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    (logs / f"{date}_{dt.datetime.now():%H%M%S}.json").write_text(
        json.dumps({"roster": people, "feedback": feedback}, indent=2, ensure_ascii=False), encoding="utf-8")

    check_excel_closed(excel)
    update_excel(excel, date, people, feedback, cfg.get("excel", "backup_dir"), dry_run=args.dry_run)
    log("Pearson VUE part done")


def run_linkedin(cfg, date, month_day, args):
    """Excel EmailId -> Outlook contact card -> LinkedIn URL -> Excel 'Linkedin Id'."""
    excel = cfg.get("excel", "path")
    log("=== LinkedIn from Outlook")
    check_excel_closed(excel)

    sheet, todo = linkedin_rows_to_do(excel, date, month_day)
    if args.limit:
        todo = todo[:args.limit]
    log(f"Sheet {sheet}: {len(todo)} email(s) without a LinkedIn Id"
        + (f" on {date:%d-%b-%Y}" if date else ""))
    if not todo:
        return
    for x in todo:
        log(f"  row {x['row']}: {x['name']} | {x['email']}")

    port = cfg.getint("chrome", "debug_port")

    def start_outlook():
        d = open_chrome(cfg, "linkedin", cfg.get("linkedin", "outlook_url"), "live.com")
        open_compose(d)
        return d, d.current_window_handle

    def back_to_outlook():
        """Return to the Outlook tab; if it was closed, open Outlook again."""
        nonlocal driver, outlook_tab
        try:
            driver.switch_to.window(outlook_tab)
            driver.execute_script("return 1")
        except Exception:
            log("    Outlook tab was closed - opening it again")
            driver, outlook_tab = start_outlook()

    driver, outlook_tab = start_outlook()

    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    progress = logs / f"linkedin_{dt.datetime.now():%Y%m%d_%H%M%S}.json"
    try:
        for i, x in enumerate(todo, 1):
            log(f"[{i}/{len(todo)}] {x['name']} | {x['email']}")
            try:
                x["linkedin"] = lookup(driver, port, outlook_tab, x["email"])
            except Exception as e:
                log(f"    error: {type(e).__name__}: {str(e).splitlines()[0][:150]}")
                back_to_outlook()
            if not x["linkedin"] and args.ask:
                x["linkedin"] = manual_lookup(driver, port, outlook_tab, x)
            log(f"    LinkedIn: {x['linkedin'] or '(not found)'}")
            try:
                close_card(driver)
            except Exception:
                back_to_outlook()
            progress.write_text(json.dumps(todo, indent=2), encoding="utf-8")
            time.sleep(1)
    except KeyboardInterrupt:
        log("Stopped - saving what was found so far")
    finally:
        try:
            discard_draft(driver)
        except Exception:
            log("Could not discard the draft - delete it from Outlook Drafts")
        found = sum(1 for x in todo if x["linkedin"])
        log(f"LinkedIn found for {found} of {len(todo)} email(s)")
        for x in todo:
            if not x["linkedin"]:
                log(f"  ! not found: {x['name']} | {x['email']}")
        if args.dry_run:
            log("Dry run - Excel file NOT saved")
        else:
            save_linkedin_results(excel, sheet, todo, cfg.get("excel", "backup_dir"))
    log("LinkedIn part done")


def run_student(cfg, date, month_day, args):
    """Excel Phone No -> Adv Pro Registration search -> Excel 'Our Student Y/N'."""
    excel = cfg.get("excel", "path")
    log("=== Our Student check in Adv Pro")
    check_excel_closed(excel)

    ws, cols, rows = month_rows(excel, date, month_day, ["our student y/n"])
    todo = [{"row": r, "name": clean_cell_text(ws.cell(r, cols["name"]).value),
             "phone": phone10(ws.cell(r, cols["phone no"]).value), "our_student": ""}
            for r in rows
            if phone10(ws.cell(r, cols["phone no"]).value) and is_blank(ws.cell(r, cols["our student y/n"]).value)]
    if args.limit:
        todo = todo[:args.limit]
    log(f"Sheet {ws.title}: {len(todo)} phone number(s) without 'Our Student Y/N'"
        + (f" on {date:%d-%b-%Y}" if date else ""))
    if not todo:
        return

    try:
        adv = AdvPro(cfg)
        adv.main.set_focus()  # keep it in front for the screen check
        for i, x in enumerate(todo, 1):
            x["our_student"] = adv.is_student(x["phone"])
            log(f"[{i}/{len(todo)}] {x['phone']} {x['name']}: {x['our_student']}")
    except KeyboardInterrupt:
        log("Stopped - saving what was checked so far")
    finally:
        updates = [{"row": x["row"], "key": ("phone no", x["phone"]),
                    "values": {"our student y/n": x["our_student"]},
                    "label": f"{x['name']} | {x['phone']} -> {x['our_student']}"}
                   for x in todo if x["our_student"]]
        if args.dry_run:
            log(f"Dry run - Excel file NOT saved ({len(updates)} row(s) checked)")
        else:
            save_cells(excel, ws.title, updates, cfg.get("excel", "backup_dir"), "'Our Student Y/N'")
    log("Our Student part done")


def run_certs(cfg, date, month_day, args):
    """Excel Linkedin Id -> LinkedIn 'Licenses & certifications' -> 'y' in RHCSA ... Anthropic."""
    excel = cfg.get("excel", "path")
    log("=== Certifications from LinkedIn")
    check_excel_closed(excel)

    ws, cols, rows = month_rows(excel, date, month_day, ["linkedin id", *CERTS])
    todo = []
    for r in rows:
        url = linkedin_profile(str(ws.cell(r, cols["linkedin id"]).value or ""))
        if url and all(is_blank(ws.cell(r, cols[c]).value) for c in CERTS):
            todo.append({"row": r, "name": clean_cell_text(ws.cell(r, cols["name"]).value),
                         "linkedin": str(ws.cell(r, cols["linkedin id"]).value).strip(),
                         "url": url, "certs": None})
    if args.limit:
        todo = todo[:args.limit]
    log(f"Sheet {ws.title}: {len(todo)} LinkedIn profile(s) not checked yet"
        + (f" on {date:%d-%b-%Y}" if date else ""))
    if not todo:
        return

    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    progress = logs / f"certs_{dt.datetime.now():%Y%m%d_%H%M%S}.json"
    delay = cfg.getfloat("certs", "delay", fallback=5)
    try:
        driver = open_chrome(cfg, "linkedin", "https://www.linkedin.com/feed/", "linkedin.com")
        for i, x in enumerate(todo, 1):
            log(f"[{i}/{len(todo)}] {x['name']} | {x['url']}")
            try:
                text = read_certifications(driver, x["url"])
            except Exception as e:
                log(f"    error: {type(e).__name__}: {str(e).splitlines()[0][:150]}")
                text = None
            if text is None:
                log("    could not open the profile - skipped")
            else:
                x["certs"] = find_certs(text)
                x["text"] = text[:4000]
                log(f"    certifications: {', '.join(c.upper() for c in x['certs']) or '(none of ours)'}")
            progress.write_text(json.dumps(todo, indent=2, ensure_ascii=False), encoding="utf-8")
            time.sleep(delay)
    except KeyboardInterrupt:
        log("Stopped - saving what was found so far")
    finally:
        updates = [{"row": x["row"], "key": ("linkedin id", x["linkedin"]),
                    "values": {c: "y" for c in x["certs"]},
                    "label": f"{x['name']} -> {', '.join(c.upper() for c in x['certs'])}"}
                   for x in todo if x["certs"]]
        if args.dry_run:
            log(f"Dry run - Excel file NOT saved ({len(updates)} profile(s) with certifications)")
        else:
            save_cells(excel, ws.title, updates, cfg.get("excel", "backup_dir"), "certifications")
    log("Certifications part done")


def run_test(cfg):
    """Offline check of the Excel step with sample data. Never saves the Excel file."""
    log("=== Offline test (sample data, Excel is NOT saved)")
    date = dt.date(2026, 10, 6)
    people = [
        {"name": "KOSAL RAJ I", "client": "Google Cloud", "code": "PR000260", "phone": "+91 9791117206"},  # already there
        {"name": "Test Person", "client": "Microsoft", "code": "AZ-900", "phone": "+91 98765 43210"},     # new
    ]
    tmp = Path(tempfile.gettempdir()) / "fb_test.csv"
    with open(tmp, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Timestamp", "Name", "Phone Number", "Email address", "Your company name",
                    "Your Designation", "Let others know about your experience"])
        w.writerow(["10/6/2026 18:05:11", "Kosal Raj", "9791117206", "kosal@example.com", "Acme", "Engineer", "Great centre"])
        w.writerow(["10/5/2026 11:00:00", "Old", "9000000000", "old@example.com", "X", "Y", "yesterday"])
    feedback = read_feedback(tmp, date, "")
    update_excel(cfg.get("excel", "path"), date, people, feedback, "backups", dry_run=True)


# ================================================================ main

MENU = """
==============================================
   Advantage Pro - Exam Automation
==============================================
   1. Pearson VUE roster + feedback  -> Excel
   2. LinkedIn from Outlook          -> Excel
   3. Our Student check in Adv Pro   -> Excel
   4. Certifications from LinkedIn   -> Excel
   5. All (1, 2, 3, 4 - same date)
   6. Test (sample data, nothing saved)
   0. Exit
"""
TASKS = {"1": "pvue", "2": "linkedin", "3": "student", "4": "certs", "5": "all", "6": "test", "0": "exit"}


def ask_task():
    print(MENU)
    while True:
        task = TASKS.get(input(">>> Choose 1-6 or 0: ").strip())
        if task:
            return task
        print("    Please type a number from the menu")


def ask_or_parse_scope(args):
    """--date / --month / ask (Enter = whole month) -> (date or None, day of the month)."""
    if args.date:
        date = parse_date(args.date)
        return date, date
    if args.month:
        return None, dt.date.fromisoformat(args.month + "-01")
    return ask_scope()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("task", nargs="?", choices=["pvue", "linkedin", "student", "certs", "all", "test"],
                    help="what to run (no task = show the menu)")
    ap.add_argument("--date", help="DD-MM-YYYY or YYYY-MM-DD (default: ask)")
    ap.add_argument("--month", help="linkedin/student/certs: YYYY-MM, every row of that month")
    ap.add_argument("--skip-roster", action="store_true", help="pvue: feedback only")
    ap.add_argument("--skip-feedback", action="store_true", help="pvue: roster only")
    ap.add_argument("--limit", type=int, help="linkedin/student/certs: only the first N rows")
    ap.add_argument("--ask", action="store_true", help="linkedin: wait for you when a profile isn't found")
    ap.add_argument("--dry-run", action="store_true", help="do everything except saving the Excel file")
    args = ap.parse_args()

    task = args.task or ask_task()
    if task == "exit":
        return
    cfg = load_config()

    if task == "test":
        run_test(cfg)

    elif task == "pvue":
        date = parse_date(args.date) if args.date else ask_date()
        run_pvue(cfg, date, args)

    elif task == "linkedin":
        run_linkedin(cfg, *ask_or_parse_scope(args), args)

    elif task == "student":
        run_student(cfg, *ask_or_parse_scope(args), args)

    elif task == "certs":
        run_certs(cfg, *ask_or_parse_scope(args), args)

    elif task == "all":
        date = parse_date(args.date) if args.date else ask_date()
        run_pvue(cfg, date, args)
        run_linkedin(cfg, date, date, args)
        run_student(cfg, date, date, args)
        run_certs(cfg, date, date, args)

    log("Done")


if __name__ == "__main__":
    main()
